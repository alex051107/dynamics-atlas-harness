"""Build the final 12-slide argument from sealed results; never grade answers here."""
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_DATA_LABEL_POSITION
r=Path(__file__).resolve().parents[1];o=r/'outputs'
assert(o/'UNBLINDING.json').exists()
g=json.loads((o/'GROUP_RESULTS.json').read_text());v=json.loads((o/'FINAL_DECISIONS.json').read_text());cost=json.loads((o/'CAMPAIGN_PROGRESS.json').read_text())
def group(e,c,a):return next(x for x in g if(x['experiment'],x['case'],x['arm'])==(e,c,a))
def m(e,c,a,k):return group(e,c,a)['median'][k]
p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
navy='20394E';blue='3674A0';amber='AF6B30';gray='536474';font='PingFang SC'
def txt(s,x,y,w,h,t,size=22,bold=False,color=navy):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True
 for j,line in enumerate(t.split('\n')):
  q=tf.paragraphs[0]if j==0 else tf.add_paragraph();q.text=line;q.font.name=font;q.font.size=Pt(size);q.font.bold=bold;q.font.color.rgb=RGBColor.from_string(color);q.space_after=Pt(15)
 return sh

def slide(title,source,note):
 s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(255,255,255)
 txt(s,.65,.5,12.05,1.05,title,28,True)
 txt(s,.68,6.78,11.9,.5,source,12,color=gray);txt(s,12.25,6.94,.4,.25,str(len(p.slides)),10,color=gray)
 s.notes_slide.notes_text_frame.text=note;return s

def chart(s,cats,series,x=.7,y=1.75,w=7.7,h=4.65,ytitle='',maximum=None):
 d=CategoryChartData();d.categories=cats
 for n,vals in series:d.add_series(n,vals)
 ch=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x),Inches(y),Inches(w),Inches(h),d).chart
 ch.has_legend=len(series)>1
 if ch.has_legend:ch.legend.position=XL_LEGEND_POSITION.BOTTOM;ch.legend.font.size=Pt(16)
 ch.category_axis.tick_labels.font.size=Pt(16);ch.category_axis.tick_labels.font.name=font;ch.value_axis.tick_labels.font.size=Pt(15)
 ch.value_axis.minimum_scale=0;ch.value_axis.tick_labels.number_format="0.#"
 ch.value_axis.has_title=bool(ytitle)
 if ytitle:ch.value_axis.axis_title.text_frame.text=ytitle;ch.value_axis.axis_title.text_frame.paragraphs[0].font.size=Pt(16)
 if maximum is not None:
  ch.value_axis.maximum_scale=maximum
  ch.value_axis.major_unit=0.5 if maximum<5 else (1 if maximum<=5 else 2)
 for se,col in zip(ch.series,[blue,amber,navy]):se.format.fill.solid();se.format.fill.fore_color.rgb=RGBColor.from_string(col)
 plot=ch.plots[0];plot.has_data_labels=True;plot.data_labels.position=XL_DATA_LABEL_POSITION.OUTSIDE_END;plot.data_labels.font.size=Pt(16);plot.data_labels.number_format='0.##'
 return ch

s=slide('Rules 的价值，要落实到具体科学回答','Dynamics Atlas · 2026-09-10 · 四层验证与三个体系的实际结果',v['summary'])
txt(s,.8,1.8,11.8,1.2,v['summary'],27,True)
txt(s,.8,3.25,11.8,2.8,'\n'.join(f"{k} {v[k]['label']}：{v[k]['short']}"for k in ['A','B','C','D']),22)
s=slide('同资料、同工具，检验辅助机制改变了什么','来源：各实验 FREEZE.json、运行回执、盲评分与费用账本','最初目的不是维持规则引擎，而是得到可核查的异质蛋白证据回答。普通组已有整理好的表。72次对照加2次ADK普通答复；每题每条件四次。评分者参与建包，条件标签隐藏，但不是独立领域专家双盲。')
txt(s,.8,1.75,11.8,.8,'研究问题：辅助机制能否让 Agent 少错、多答对？',26,True)
txt(s,.8,2.9,6.2,3.1,'准入卡：缺陷识别与结论降级\n一次反馈：错误、遗漏与新增错误\n方法卡：适用性与科学正确\n拆题：有证据的子问题覆盖',23)
txt(s,7.3,2.9,5.1,3.1,f"Luna · medium\n每组每题 4 次 · 新上下文\n新增费用 ${cost['total_settled_new_cost_usd']:.4f}\n暴露开发材料 + 1 个新ADK来源",22)
s=slide('项目已从规则覆盖转向科学结果与受控比较','来源：9月8日报告与战略审查；9月9–10日冻结材料、原始运行记录','规则33条是来源材料。Q05暴露了关键分析由开发者安排。HSP90原生NOE对照是跨来源结果；DHFR、ADK主要是MD内部描述。本轮对组件逐项比较，不把工程可运行当作效果。')
for y,a,b in [(1.75,'整理 33 条候选规则','形成有来源的前提与边界，但未测独立增量'),(2.85,'Q05 开发与 9月8日重置','科学问题和分析协议居中，停止按规则数量推进'),(3.95,'HSP90 → DHFR → ADK','得到具体数值，也发现列含义和周期表示问题'),(5.05,'本轮四层比较','按预先冻结的标准决定保留、简化或暂不投入')]:txt(s,.8,y,4.2,.65,a,23,True);txt(s,5.1,y,7.1,.8,b,21)
s=slide('HSP90 的开放方向不等于接近开放参照','Henot等，2022，doi:10.1038/s41467-022-35399-8；项目原生NOE逐轨迹对照','闭合起始20条有10条出现持续开放方向，分成5条起始已有、5条后来出现、10条从未出现。图中闭合起始只展示这10条，与全部20条开放起始对照比较。容差1Å；原生NOE第二列与伪距离第三列不是同一量。')
chart(s,['仅相对方向','部分一致','一致'],[('闭合起始中趋开的10条',[9,1,0]),('开放起始20条',[1,1,18])],ytitle='轨迹条数',maximum=20)
txt(s,8.85,1.95,3.65,3.9,'10 = 5 起始就开\n       + 5 后来趋开\n\n1 Å 下仅少数接近参照\n方向、绝对读数分别回答',22)
s=slide('DHFR 的局部距离差异要先通过周期核对','Cetin等，2023，PMC10428214，Fig.4；项目修正表与VMD核对','展示M20-O3P平均最近周期配体副本距离。每条件一条轨迹，990帧。旧两份Agent均未识别包装坐标缺陷；修正数字属于开发者。距离更小不能单独证明氢键、抑制机制或解离速率。')
chart(s,['WT','L28R'],[('TMP',[8.687677,10.435783]),("4′-DTMP",[4.622056,4.808910])],ytitle='M20–O3P 平均距离 / Å')
txt(s,8.85,1.95,3.65,4.1,'两条件都更近\n\n每条件仅 1 条轨迹\n旧原答未识别坏表\n修正数字不归功于 Agent',22)
s=slide('ADK 两条轨迹未显示两个结构域共同闭合','Orädd等，2021，Sci. Adv. 7:eabi5514；Zenodo 5583119；项目独立复算','完整分子直接距离，与GROMACS全帧交叉核对差约0.005Å。域定义为项目预定的Cα对平均距离；正值表示远离。首末各10%帧，开放窗口44.8ns、闭合窗口33.2ns。apo模拟不是ATP光释放实验；不能把4.3ms实验中间态直接映射为轨迹时钟。')
chart(s,['开态起始','闭态起始'],[('NMP–CORE',[.0707211184,1.0908948441]),('LID–CORE',[2.2904106995,1.1764212186])],ytitle='末窗减首窗 / Å')
txt(s,8.85,1.95,3.65,4.2,'四项变化均为正\n分布仍有重叠\n\n合法大分子距离\n不因超过半盒长而拒收',22)
s=slide('准入层必须把可信参照与自动检查分开','来源：E1a_RESULTS、E1A_CHECK_COVERAGE_AUDIT、E1b盲评','6/6是与可信参照一致性测试，0/3干净包与参照完全相同。独立坐标重算是上游工作。图显示实际缺陷识别与降级行为，四次每组。'+v['B']['reason'])
chart(s,['缺陷识别','据此降级'],[('无卡',[group('E1b','Q-D-defect','D')['sum'][k]for k in ['defect_detected','defect_downgraded']]),('有卡',[group('E1b','Q-D-defect','CARD')['sum'][k]for k in ['defect_detected','defect_downgraded']])],ytitle='4次中的次数',maximum=4)
txt(s,8.85,1.95,3.65,4.1,'6/6 植入缺陷检出\n0/3 自一致性误报\n\n'+v['B']['short'],22)
s=slide('结论检查要按真实纠错与误报评价','来源：E2逐题盲评、首次/最终提交；E0历史回放','历史回放当前检查器未检出7处目标核心错误，唯一提示把Å^-1指数误认数字。E2三题缺PUBLIC_FACTS，角色检查未被测试；实际反馈另做语义核对。'+v['D']['reason'])
chart(s,['HSP90','DHFR','ADK'],[('无反馈',[m('E2',c,'D','core_errors')for c in ['Q-H2','Q-D-fixed','Q-A1']]),('一次反馈',[m('E2',c,'FEEDBACK','core_errors')for c in ['Q-H2','Q-D-fixed','Q-A1']])],ytitle='核心错误中位数',maximum=max(1,max(x['median']['core_errors']for x in g if x['experiment']=='E2')+1))
txt(s,8.85,1.95,3.65,4.2,v['D']['short']+'\n\n数字出现过\n不等于计算有依据\n同时检查遗漏与过强结论',22)
s=slide('方法族匹配尚不能代替操作适用性判断','来源：E3_CARD_RETRIEVAL、冻结卡片与E3逐份盲评','HSP90取到1张BME卡，但当前NOE违例对照不要求重加权。Q05取到4张；DHFR/ADK的MD比较取到0张。主判据仅四个冻结术语，额外BME问题为探索性分析。'+v['C']['reason'])
chart(s,['HSP90','Q05'],[(label,[m('E3',c,a,'core_correct')for c in ['Q-H2','Q05']])for a,label in [('P','七条协议'),('R_old','整段规则'),('R_scoped','方法卡')]],ytitle='5个核心单元：正确中位数',maximum=5)
txt(s,8.85,1.95,3.65,4.1,v['C']['short']+'\n\nHSP90：1 张\nQ05：4 张\n检索命中不等于适用',22)
s=slide('拆题比较衡量明确需求后多回答了什么','来源：E4冻结拆题、共有评分依据和逐题盲评','两版资料工具相同。拆题版明确要求更多子问题，估计对象是框题包，不是严格等长同义改写，也不是自主问题发现。冻结按2/2题子问题覆盖改善且过强结论不增。'+v['A']['reason'])
chart(s,['HSP90','ADK'],[('原题',[m('E4',c,'ORIGINAL','subquestion_coverage')for c in ['HSP90_Q01','Q-A1']]),('拆题',[m('E4',c,'SPLIT','subquestion_coverage')for c in ['HSP90_Q01','Q-A1']])],ytitle='3个子问题：覆盖中位数',maximum=3)
txt(s,8.85,1.95,3.65,4.1,v['A']['short']+'\n\n覆盖与核心错误分列\n一个ADK来源\n不能推算总体正确率',22)
s=slide('所有判断都能回到论文、原答和计算记录','报告与证据：github.com/alex051107/dynamics-atlas-harness/pull/27','完整来源定位在科学读本、claim_source_map和逐份评分。论文结论与本地计算、Agent实际分析分别记录。Bengtsen和Fuertes卡片来源及具体限制见METHOD_CARD_DESIGN_ANALYSIS_ZH。')
txt(s,.85,1.7,11.7,4.75,'Henot等（2022）· HSP90原生NOE与模拟\ndoi:10.1038/s41467-022-35399-8\nCetin等（2023）· DHFR抑制剂比较\nPMC10428214，Table 1 / Fig.3 / Fig.4\nOrädd等（2021）· Tracking the ATP-binding response in adenylate kinase in real time\ndoi:10.1126/sciadv.abi5514；Zenodo 5583119',21)
s=slide('下一步只保留得到实际支持的组件','完整报告、原答、评分与费用：PR #27 · review/four-layer-20260910',v['next_action'])
txt(s,.85,1.8,11.6,1.2,v['summary'],27,True)
txt(s,.85,3.15,11.6,2.9,v['next_action'],24)
p.save(o/'COLLABORATOR_REPORT.pptx');print('Saved 12-slide final deck from sealed results')
